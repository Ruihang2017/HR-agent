from shortlist.pipeline.extract import extract_text


def test_txt_extraction(tmp_path):
    p = tmp_path / "resume.txt"
    p.write_text("Alex Chen\nBarista", encoding="utf-8")
    result = extract_text(p)
    assert result.error is None
    assert "Barista" in result.text


def test_docx_extraction(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Alex Chen")
    doc.add_paragraph("Barista at Beans & Co")
    p = tmp_path / "resume.docx"
    doc.save(str(p))
    result = extract_text(p)
    assert result.error is None
    assert "Beans & Co" in result.text


def test_corrupt_pdf_degrades_gracefully(tmp_path):
    p = tmp_path / "resume.pdf"
    p.write_bytes(b"%PDF-1.4 this is not really a pdf")
    result = extract_text(p)
    assert result.text is None
    assert result.error is not None


def test_unsupported_extension(tmp_path):
    p = tmp_path / "resume.pages"
    p.write_text("hi", encoding="utf-8")
    result = extract_text(p)
    assert result.text is None
    assert "unsupported" in result.error
